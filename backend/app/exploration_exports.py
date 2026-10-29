"""Exploratory report exporters.

Builds a human-readable XLSX or a lossless JSON from an exploratory run's
``ExplorationResponse`` (the persisted report, which carries its metrics), or
from every run of a job. Like ``abcd_exports`` this is pure over already-loaded
data: routes do the reads and hand the models in, so it is unit-testable
without a database.

A report is described once, as a list of ``_Section`` tables. A single run puts
each section on its own sheet; a job puts all of one video's sections on one
sheet, stacked, next to a cross-video Summary sheet (mirroring the ABCD job
export).
"""

from __future__ import annotations

import json
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from app.exports import (
    XLSX_CONTENT_TYPE,
    autosize,
    cell_value,
    safe_sheet_title,
    workbook_bytes,
)
from app.models import ExplorationResponse
from clip_scribe.parser.explore.models import ExploreReport, TimeSpan

VALID_FORMATS = frozenset({"json", "xlsx"})
CONTENT_TYPES = {
    "json": "application/json",
    "xlsx": XLSX_CONTENT_TYPE,
}

_HEADER_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=13)
# Prose columns wrap and are capped so one long summary can't blow out a sheet.
_WRAP = Alignment(wrap_text=True, vertical="top")
_PROSE_WIDTH = 90
# Excel's hard limit on characters in one cell.
_MAX_CELL_CHARS = 32767


@dataclass(frozen=True)
class RunExploration:
    """One run's data for a job-level export."""

    video_name: str
    response: ExplorationResponse
    # The run's raw speech segments (empty when the run has no speech).
    audio_segments: Sequence[Mapping[str, Any]] = ()


@dataclass(frozen=True)
class _Section:
    """One table of a report: a sheet in a run export, a block in a job sheet."""

    title: str
    headers: Sequence[str]
    rows: Sequence[Sequence[object]]
    # 1-based columns holding prose, wrapped and width-capped.
    wrap_columns: frozenset[int] = frozenset()


def full_transcript(segments: Iterable[Mapping[str, Any]]) -> str:
    """All speech segments joined into one string, cut to fit an Excel cell."""
    text = " ".join(s["text"].strip() for s in segments if s["text"].strip())
    if len(text) > _MAX_CELL_CHARS:
        marker = " [truncated]"
        text = text[: _MAX_CELL_CHARS - len(marker)] + marker
    return text


def run_json(response: ExplorationResponse) -> bytes:
    return response.model_dump_json(indent=2).encode("utf-8")


def job_json(runs: Sequence[RunExploration]) -> bytes:
    """Every run's report in one lossless document, tagged with its video."""
    body = {
        "runs": [
            {"video_name": r.video_name, **r.response.model_dump(mode="json")}
            for r in runs
        ]
    }
    return json.dumps(body, indent=2, ensure_ascii=False).encode("utf-8")


def run_xlsx(
    video_name: str,
    response: ExplorationResponse,
    audio_segments: Sequence[Mapping[str, Any]] = (),
) -> bytes:
    """One run, one sheet per section: Overview, Chapters, Transcript, On-screen
    text, Visuals, and Findings (only when the findings stage produced any).

    ``audio_segments`` are the run's raw speech segments; their text is joined
    into the Transcript sheet's full-transcript row.
    """
    sections = _sections(video_name, response, full_transcript(audio_segments))
    wb = Workbook()
    first = wb.active
    assert first is not None
    for i, section in enumerate(sections):
        ws = first if i == 0 else wb.create_sheet()
        ws.title = section.title
        _table(ws, section)
    return workbook_bytes(wb)


def job_xlsx(runs: Sequence[RunExploration]) -> bytes:
    """Whole job: a Summary sheet across videos, then one sheet per video."""
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    summary.title = "Summary"
    _table(summary, _job_summary(runs))
    used = {"summary"}
    for i, run in enumerate(runs, start=1):
        sections = _sections(
            run.video_name, run.response, full_transcript(run.audio_segments)
        )
        title = safe_sheet_title(run.video_name, i, used)
        _stacked(wb.create_sheet(title), sections)
    return workbook_bytes(wb)


# ── layout helpers ──────────────────────────────────────────────────────────


def _table(ws: Worksheet, section: _Section) -> None:
    """The section as a whole sheet: header row, frozen, wrapped prose."""
    ws.append(list(section.headers))
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    for row in section.rows:
        ws.append([cell_value(v) for v in row])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in section.wrap_columns:
                cell.alignment = _WRAP
    ws.freeze_panes = "A2"
    autosize(ws, column_max_widths={c: _PROSE_WIDTH for c in section.wrap_columns})


def _stacked(ws: Worksheet, sections: Sequence[_Section]) -> None:
    """All sections on one sheet: a title row, the table, then a blank row."""
    wrapped: set[int] = set()
    for i, section in enumerate(sections):
        if i:
            ws.append([])
        ws.append([section.title])
        ws.cell(row=ws.max_row, column=1).font = _TITLE_FONT
        ws.append(list(section.headers))
        for cell in ws[ws.max_row]:
            cell.font = _HEADER_FONT
        for row in section.rows:
            ws.append([cell_value(v) for v in row])
        wrapped |= section.wrap_columns
    for row in ws.iter_rows():
        for cell in row:
            if cell.column in wrapped:
                cell.alignment = _WRAP
    autosize(ws, column_max_widths={c: _PROSE_WIDTH for c in wrapped})


def _clock(sec: float) -> str:
    minutes, seconds = divmod(int(sec), 60)
    return f"{minutes}:{seconds:02d}"


def _spans(spans: list[TimeSpan]) -> str:
    return ", ".join(f"{_clock(s.start_sec)}-{_clock(s.end_sec)}" for s in spans)


def _percent(fraction: float | None) -> str | None:
    return None if fraction is None else f"{fraction:.0%}"


# ── report sections ─────────────────────────────────────────────────────────


def _sections(
    video_name: str, response: ExplorationResponse, transcript: str
) -> list[_Section]:
    report = response.report
    if report is None:
        raise ValueError("run has no exploration report to export")
    sections = [
        _overview(video_name, response, report),
        _chapters(report),
        _transcript(report, transcript),
        _on_screen_text(report),
        _visuals(report),
    ]
    if report.findings:
        sections.append(_findings(report))
    return sections


def _job_summary(runs: Sequence[RunExploration]) -> _Section:
    """One row per video: the headline numbers and the overview summary."""
    rows: list[Sequence[object]] = []
    for run in runs:
        report = run.response.report
        assert report is not None
        m = report.metrics
        rows.append(
            (
                run.video_name,
                m.duration_sec,
                m.shot_count,
                m.cuts_per_minute,
                _percent(m.speech_coverage),
                _percent(m.text_coverage),
                len(report.chapters),
                report.summary,
            )
        )
    return _Section(
        "Summary",
        (
            "Video",
            "Duration (s)",
            "Shots",
            "Cuts per minute",
            "Speech coverage",
            "On-screen text coverage",
            "Chapters",
            "Summary",
        ),
        rows,
        frozenset({8}),
    )


def _overview(
    video_name: str, response: ExplorationResponse, report: ExploreReport
) -> _Section:
    m = report.metrics
    rows: list[tuple[str, object]] = [
        ("Video", video_name),
        ("Run ID", response.run_id),
        ("Model", response.model),
        ("Created", response.created_at),
        ("Summary", report.summary),
        ("Duration (s)", m.duration_sec),
        ("Shots", m.shot_count),
        ("Average shot (s)", m.avg_shot_sec),
        ("Cuts per minute", m.cuts_per_minute),
        ("Speech coverage", _percent(m.speech_coverage)),
        ("On-screen text coverage", _percent(m.text_coverage)),
        ("Words per minute", m.words_per_minute),
        (
            "Face coverage",
            _percent(m.face_coverage) if m.objects_tracked else "Not collected",
        ),
    ]
    return _Section("Overview", ("Field", "Value"), rows, frozenset({2}))


def _chapters(report: ExploreReport) -> _Section:
    return _Section(
        "Chapters",
        ("#", "Title", "Start", "End", "Summary"),
        [
            (i, c.title, _clock(c.start_sec), _clock(c.end_sec), c.summary)
            for i, c in enumerate(report.chapters, start=1)
        ],
        frozenset({5}),
    )


def _transcript(report: ExploreReport, transcript: str) -> _Section:
    """The full speech first, then the digest (summary, topics, key claims).

    The full text is shown even when the digest failed; the status row only
    appears when there is neither.
    """
    digest = report.transcript
    rows: list[tuple[str, object, object]] = []
    if transcript:
        rows.append(("Full transcript", None, transcript))
    if digest is not None:
        rows.append(("Summary", None, digest.summary))
        rows.extend(("Topic", None, topic) for topic in digest.topics)
        rows.extend(
            ("Key claim", _clock(c.start_sec), c.text) for c in digest.key_claims
        )
    if not rows:
        rows.append(("Status", None, "No speech in this video"))
    return _Section("Transcript", ("Section", "Time", "Text"), rows, frozenset({3}))


def _on_screen_text(report: ExploreReport) -> _Section:
    digest = report.on_screen_text
    headers = ("Section", "Text", "Raw OCR reading", "Category", "Visible", "Seconds")
    if digest is None:
        rows: list[Sequence[object]] = [
            ("Status", "No on-screen text detected", None, None, None, None)
        ]
    else:
        rows = [
            ("Summary", digest.summary, None, None, None, None),
            *(
                (
                    "Card",
                    c.clean_text,
                    c.text,
                    c.category.value,
                    _spans(c.spans),
                    c.total_sec,
                )
                for c in digest.cards
            ),
        ]
    return _Section("On-screen text", headers, rows, frozenset({2}))


def _visuals(report: ExploreReport) -> _Section:
    digest = report.visuals
    headers = ("Section", "Text / label", "Prominence", "Screen time (s)", "First seen")
    if digest is None:
        rows: list[Sequence[object]] = [("Status", "No visual data", None, None, None)]
    else:
        rows = [
            ("Summary", digest.summary, None, None, None),
            *(
                (
                    "Subject",
                    s.label,
                    s.prominence,
                    s.screen_time_sec,
                    _clock(s.first_seen_sec),
                )
                for s in digest.subjects
            ),
        ]
    return _Section("Visuals", headers, rows, frozenset({2}))


def _findings(report: ExploreReport) -> _Section:
    return _Section(
        "Findings",
        ("#", "Observation", "Evidence"),
        [
            (i, f.observation, ", ".join(_clock(s) for s in f.evidence_secs))
            for i, f in enumerate(report.findings, start=1)
        ],
        frozenset({2}),
    )

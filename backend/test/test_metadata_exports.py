"""Focused tests for the pure video-metadata export builders."""

import io
import json
from typing import cast

import openpyxl

from app import metadata_exports
from clip_scribe.extractor.extractor_models import ExtractionSummary


def _summary() -> ExtractionSummary:
    return cast(
        ExtractionSummary,
        {
            "global_stats": {
                "total_shots": 1,
                "video_duration": 4.0,
                "avg_shot_duration": 4.0,
                "dynamic_start": {
                    "detected": False,
                    "first_shot_duration": 4.0,
                    "criteria": "First shot < 3.0s",
                },
                "quick_pacing_intro": {
                    "detected": False,
                    "shot_count": 1,
                    "shots": [0],
                    "criteria": ">= 5 shots in first 5s",
                },
                "quick_pacing_general": {
                    "detected": True,
                    "rapid_fire_segments": [
                        {
                            "start_time": 0.0,
                            "end_time": 4.0,
                            "duration": 4.0,
                            "shot_count": 1,
                            "shot_indices": [0],
                        }
                    ],
                    "criteria": ">= 5 shots in any 5s window",
                },
            },
            "visual_objects": [
                {
                    "global_id": 7,
                    "label": "+product",
                    "occurrences": [
                        {
                            "shot_index": 0,
                            "lifespan": [0.5, 2.5],
                            "screen_coverage": 0.2,
                            "velocity_px_sec": 12.0,
                            "growth_factor": 1.2,
                            "direction": "right",
                            "centrality_score": 0.1,
                            "screen_time_ratio": 0.5,
                            "quadrant": "center",
                        }
                    ],
                }
            ],
            "text_events": [{"second": 1, "text": ["=headline", "second line"]}],
            "audio_segments": [
                {
                    "start": 0.0,
                    "end": 2.0,
                    "text": "@spoken words",
                    "confidence": 0.8,
                }
            ],
            "scene_descriptions": [
                {
                    "shot_index": 0,
                    "start_time": 0.0,
                    "end_time": 4.0,
                    "description": "=scene description",
                }
            ],
            "shot_boundaries": [
                {"index": 0, "start": 0.0, "end": 4.0, "duration": 4.0}
            ],
            "frame_detections": [
                {
                    "shot_index": 0,
                    "frame_idx": 10,
                    "timestamp_sec": 1.0,
                    "source": "sam_mask",
                    "label": "raw-frame-marker",
                    "text": None,
                    "box_x1": 1.0,
                    "box_y1": 2.0,
                    "box_x2": 3.0,
                    "box_y2": 4.0,
                    "confidence": 0.9,
                    "object_id": 7,
                }
            ],
        },
    )


def _report(**kwargs: object) -> metadata_exports.RunMetadataReport:
    values = {
        "run_id": "run-1",
        "video_name": "=video.mp4",
        "video_type": "ad",
        "created_at": "2026-01-01T00:00:00Z",
        "extraction_mode": "with_objects",
        "summary": _summary(),
        **kwargs,
    }
    return metadata_exports.RunMetadataReport(**values)  # type: ignore[arg-type]


def test_run_json_is_lossless_and_includes_frame_detections():
    payload = json.loads(metadata_exports.run_json(_report()))

    assert payload["schema_version"] == 1
    assert payload["exported_at"].endswith("Z")
    assert payload["run"]["extraction_mode"] == "with_objects"
    detections = payload["extraction_summary"]["frame_detections"]
    assert detections[0]["label"] == "raw-frame-marker"


def test_run_xlsx_has_human_sections_and_escapes_formula_text():
    workbook = openpyxl.load_workbook(io.BytesIO(metadata_exports.run_xlsx(_report())))

    assert workbook.sheetnames == [
        "Overview",
        "Shots & Scenes",
        "Object Summary",
        "Object Occurrences",
        "On-screen Text",
        "Transcript",
        "Pacing",
        "Definitions",
    ]
    assert workbook["Overview"]["B2"].value == "'=video.mp4"
    assert workbook["Shots & Scenes"]["E2"].value == "'=scene description"
    assert workbook["Object Summary"]["B2"].value == "'+product"
    assert workbook["On-screen Text"]["C2"].value.startswith("'=headline")
    assert workbook["Transcript"]["E2"].value == "'@spoken words"
    definitions = {
        row[0].value: row[1].value
        for row in workbook["Definitions"].iter_rows(min_row=2)
    }
    assert "JSON only" in definitions["Raw frame detections"]
    assert "0 is centred" in definitions["Distance from centre"]
    assert not any(
        cell.value == "raw-frame-marker"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )


def test_no_object_workbook_marks_data_as_not_collected():
    summary = _summary()
    summary["visual_objects"] = []
    report = _report(summary=summary, extraction_mode="no_objects")
    workbook = openpyxl.load_workbook(io.BytesIO(metadata_exports.run_xlsx(report)))

    overview = {row[0].value: row[1].value for row in workbook["Overview"].iter_rows()}
    assert overview["Object tracking"] == "Not performed"
    assert overview["Object identities"] == "Not collected"
    assert (
        workbook["Object Summary"]["A2"].value
        == "Object tracking was not performed for this run."
    )


def test_job_exports_keep_incomplete_children_visible():
    envelope = metadata_exports.JobMetadataEnvelope(
        job_id="batch-1",
        job_name="Campaign",
        entries=(
            metadata_exports.BatchRunEntry(
                job_id="child-1",
                run_id="run-1",
                video_name="=video.mp4",
                status="completed",
                report=_report(),
            ),
            metadata_exports.BatchRunEntry(
                job_id="child-2",
                run_id=None,
                video_name="failed.mp4",
                status="failed",
            ),
        ),
    )

    workbook = openpyxl.load_workbook(io.BytesIO(metadata_exports.job_xlsx(envelope)))
    assert workbook["Summary"]["A2"].value == "'=video.mp4"
    assert workbook["Summary"]["D2"].value == "Yes"
    assert workbook["Summary"]["A3"].value == "failed.mp4"
    assert workbook["Summary"]["D3"].value == "No"
    assert workbook["Shots & Scenes"]["A2"].value == "'=video.mp4"
    assert workbook["Global Stats"]["K2"].value == "0"
    assert workbook["Definitions"]["A2"].value == "Raw frame detections"

    payload = json.loads(metadata_exports.job_json(envelope))
    assert payload["exported_at"].endswith("Z")
    assert (
        payload["runs"][0]["metadata"]["extraction_summary"]["frame_detections"][0][
            "label"
        ]
        == "raw-frame-marker"
    )
    assert payload["runs"][1]["metadata"] is None

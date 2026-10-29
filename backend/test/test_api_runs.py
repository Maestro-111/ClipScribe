"""API tests for the read-only /runs/* endpoints (web-app-plan §5, §11).

Seeds a run and its child rows in a temp SQLite DB, overrides get_reader, and
asserts response shapes plus 404 on an unknown run.
"""

import os

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine

from app import settings as settings_mod
from app.deps import current_user_id, get_reader
from app.main import app
from clip_scribe.db.writer import ClipScribeWriterDB
from clip_scribe.db.schema import (
    metadata_obj,
    runs_table,
    global_stats_table,
    shot_boundaries_table,
    visual_object_occurrences_table,
    text_events_table,
    audio_segments_table,
    scene_descriptions_table,
    frame_detections_table,
    abcd_results_table,
    jobs_table,
)

RUN_ID = "r1"


@pytest.fixture
def client(tmp_path, async_db_adapters):
    os.environ["CLIPSCRIBE_API_LOAD_MODELS"] = "false"
    settings_mod.get_settings.cache_clear()

    database_url = f"sqlite:///{tmp_path / 'runs.db'}"
    engine = create_engine(database_url)
    metadata_obj.create_all(engine)
    with engine.begin() as conn:
        # A run is only reachable through the job that produced it; seed that
        # owning job so run_owner() resolves to "local" (the test user) and the
        # ownership guard admits the read.
        conn.execute(
            jobs_table.insert(),
            {
                "job_id": "j1",
                "run_id": RUN_ID,
                "status": "completed",
                "created_by": "local",
            },
        )
        conn.execute(runs_table.insert(), {"run_id": RUN_ID, "video_name": "ad.mp4"})
        conn.execute(
            global_stats_table.insert(),
            {"run_id": RUN_ID, "total_shots": 2, "video_duration": 3.0},
        )
        conn.execute(
            shot_boundaries_table.insert(),
            [
                {
                    "run_id": RUN_ID,
                    "shot_index": 0,
                    "start_sec": 0.0,
                    "end_sec": 1.5,
                    "duration_sec": 1.5,
                },
                {
                    "run_id": RUN_ID,
                    "shot_index": 1,
                    "start_sec": 1.5,
                    "end_sec": 3.0,
                    "duration_sec": 1.5,
                },
            ],
        )
        conn.execute(
            visual_object_occurrences_table.insert(),
            {
                "run_id": RUN_ID,
                "global_id": 1,
                "label": "car",
                "shot_index": 0,
                "lifespan_start": 0.0,
                "lifespan_end": 1.0,
            },
        )
        conn.execute(
            text_events_table.insert(),
            {"run_id": RUN_ID, "second": 0, "line_index": 0, "text": "RAM"},
        )
        conn.execute(
            audio_segments_table.insert(),
            {
                "run_id": RUN_ID,
                "start_time": 0.0,
                "end_time": 1.0,
                "text": "hi",
                "confidence": 0.9,
            },
        )
        conn.execute(
            scene_descriptions_table.insert(),
            {
                "run_id": RUN_ID,
                "shot_index": 0,
                "start_time": 0.0,
                "end_time": 1.5,
                "description": "a car",
            },
        )
        conn.execute(
            frame_detections_table.insert(),
            [
                {
                    "run_id": RUN_ID,
                    "frame_idx": 0,
                    "timestamp_sec": 0.0,
                    "source": "dino",
                    "box_x1": 0,
                    "box_y1": 0,
                    "box_x2": 1,
                    "box_y2": 1,
                },
                {
                    "run_id": RUN_ID,
                    "frame_idx": 60,
                    "timestamp_sec": 2.0,
                    "source": "ocr",
                    "box_x1": 0,
                    "box_y1": 0,
                    "box_x2": 1,
                    "box_y2": 1,
                },
            ],
        )
        conn.execute(
            abcd_results_table.insert(),
            {
                "run_id": RUN_ID,
                "platform": "youtube",
                "feature_name": "logo",
                "evaluation": True,
            },
        )

    async_reader, _ = async_db_adapters(database_url)
    app.dependency_overrides[get_reader] = lambda: async_reader
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def test_get_run(client):
    r = client.get(f"/runs/{RUN_ID}")
    assert r.status_code == 200
    assert r.json()["video_name"] == "ad.mp4"


def test_get_run_404(client):
    r = client.get("/runs/missing")
    assert r.status_code == 404
    assert r.headers["content-type"] == "application/problem+json"


def test_run_reads_hide_another_users_run(client):
    app.dependency_overrides[current_user_id] = lambda: "other"

    assert client.get(f"/runs/{RUN_ID}").status_code == 404
    assert client.get(f"/runs/{RUN_ID}/objects").status_code == 404


def test_global_stats_bundles_shot_boundaries(client):
    r = client.get(f"/runs/{RUN_ID}/global-stats")
    assert r.status_code == 200
    body = r.json()
    assert body["global_stats"]["total_shots"] == 2
    assert [s["shot_index"] for s in body["shot_boundaries"]] == [0, 1]


def test_objects_text_audio_scenes(client):
    assert client.get(f"/runs/{RUN_ID}/objects").json()[0]["label"] == "car"
    assert client.get(f"/runs/{RUN_ID}/text-events").json()[0]["text"] == "RAM"
    assert client.get(f"/runs/{RUN_ID}/audio-segments").json()[0]["text"] == "hi"
    assert client.get(f"/runs/{RUN_ID}/scenes").json()[0]["description"] == "a car"


def test_frames_window(client):
    assert len(client.get(f"/runs/{RUN_ID}/frames").json()) == 2
    windowed = client.get(
        f"/runs/{RUN_ID}/frames", params={"from": 1.0, "to": 3.0}
    ).json()
    assert [f["frame_idx"] for f in windowed] == [60]


def test_parser(client):
    body = client.get(f"/runs/{RUN_ID}/abcd").json()
    assert body[0]["feature_name"] == "logo"


def test_subresource_404_on_unknown_run(client):
    assert client.get("/runs/missing/objects").status_code == 404
    assert client.get("/runs/missing/global-stats").status_code == 404


def test_run_exposes_parsing_mode(client):
    assert client.get(f"/runs/{RUN_ID}").json()["parsing_mode"] == "abcd"


def test_abcd_routes_409_on_exploratory_run(client, tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'runs.db'}")
    with engine.begin() as conn:
        conn.execute(
            jobs_table.update()
            .where(jobs_table.c.run_id == RUN_ID)
            .values(parsing_mode="exploratory")
        )

    assert client.get(f"/runs/{RUN_ID}").json()["parsing_mode"] == "exploratory"
    assert client.get(f"/runs/{RUN_ID}/abcd").status_code == 409
    assert client.get(f"/runs/{RUN_ID}/abcd/export").status_code == 409
    assert (
        client.post(f"/runs/{RUN_ID}/abcd/chat", json={"message": "hi"}).status_code
        == 409
    )
    # Extraction reads stay available for exploratory runs.
    assert client.get(f"/runs/{RUN_ID}/objects").status_code == 200


def test_exploration_route(client, tmp_path):
    # ABCD run: the exploration endpoint is not available.
    assert client.get(f"/runs/{RUN_ID}/exploration").status_code == 409
    assert (
        client.post(
            f"/runs/{RUN_ID}/exploration/chat", json={"message": "hi"}
        ).status_code
        == 409
    )

    engine = create_engine(f"sqlite:///{tmp_path / 'runs.db'}")
    with engine.begin() as conn:
        conn.execute(
            jobs_table.update()
            .where(jobs_table.c.run_id == RUN_ID)
            .values(parsing_mode="exploratory")
        )

    # No report persisted yet.
    assert client.get(f"/runs/{RUN_ID}/exploration").json()["report"] is None

    ClipScribeWriterDB(engine=engine).save_exploration_result(
        RUN_ID, 1, _explore_report(), "gpt-test"
    )
    body = client.get(f"/runs/{RUN_ID}/exploration").json()
    assert body["report"]["summary"] == "an ad"
    # Metrics are part of the stored report.
    assert body["report"]["metrics"]["duration_sec"] == 3.0
    assert body["model"] == "gpt-test"


def _explore_report(**overrides):
    report = {
        "schema_version": 1,
        "summary": "an ad",
        "chapters": [
            {"title": "Open", "start_sec": 0.0, "end_sec": 1.5, "summary": "start"}
        ],
        "transcript": {
            "summary": "talks",
            "topics": ["cars"],
            "key_claims": [{"text": "claim A", "start_sec": 1.0}],
        },
        "on_screen_text": {
            "summary": "branding",
            "cards": [
                {
                    "text": "JEEP.",
                    "spans": [{"start_sec": 0.0, "end_sec": 2.0}],
                    "total_sec": 2.0,
                    "clean_text": "Jeep",
                    "category": "brand",
                }
            ],
        },
        "visuals": {"summary": "a car", "subjects": []},
        "findings": [],
        "metrics": {
            "duration_sec": 3.0,
            "shot_count": 2,
            "avg_shot_sec": 1.5,
            "cuts_per_minute": 20.0,
            "speech_coverage": 0.5,
            "text_coverage": 0.25,
            "words_per_minute": 120.0,
            "face_coverage": None,
            "objects_tracked": False,
            "objects": [],
        },
    }
    report.update(overrides)
    return report


def test_exploration_export(client, tmp_path):
    import io

    from openpyxl import load_workbook

    engine = create_engine(f"sqlite:///{tmp_path / 'runs.db'}")
    with engine.begin() as conn:
        conn.execute(
            jobs_table.update()
            .where(jobs_table.c.run_id == RUN_ID)
            .values(parsing_mode="exploratory")
        )

    # No report yet: nothing to export.
    assert client.get(f"/runs/{RUN_ID}/exploration/export").status_code == 409

    ClipScribeWriterDB(engine=engine).save_exploration_result(
        RUN_ID, 1, _explore_report(), "gpt-test"
    )

    resp = client.get(f"/runs/{RUN_ID}/exploration/export?format=xlsx")
    assert resp.status_code == 200
    assert "ad.mp4_exploration.xlsx" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    # No findings were produced, so no Findings sheet.
    assert wb.sheetnames == [
        "Overview",
        "Chapters",
        "Transcript",
        "On-screen text",
        "Visuals",
    ]
    overview = {
        row[0].value: row[1].value for row in wb["Overview"].iter_rows(min_row=2)
    }
    assert overview["Summary"] == "an ad"
    # The raw speech belongs on the Transcript sheet, not the overview.
    assert "Transcript" not in overview
    transcript = {
        row[0]: row[2]
        for row in wb["Transcript"].iter_rows(min_row=2, values_only=True)
    }
    # Full speech from the extraction data, ahead of the digest.
    assert transcript["Full transcript"] == "hi"
    assert transcript["Summary"] == "talks"
    assert overview["Duration (s)"] == 3.0
    # Objects weren't tracked: faces were not collected, not "0%".
    assert overview["Face coverage"] == "Not collected"
    card = [
        r
        for r in wb["On-screen text"].iter_rows(min_row=2, values_only=True)
        if r[0] == "Card"
    ][0]
    assert card[1:5] == ("Jeep", "JEEP.", "brand", "0:00-0:02")

    body = client.get(f"/runs/{RUN_ID}/exploration/export?format=json").json()
    assert body["report"]["summary"] == "an ad"
    assert (
        client.get(f"/runs/{RUN_ID}/exploration/export?format=pdf").status_code == 400
    )


def test_exploration_export_409_on_abcd_run(client):
    assert client.get(f"/runs/{RUN_ID}/exploration/export").status_code == 409


def _make_exploratory(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'runs.db'}")
    with engine.begin() as conn:
        conn.execute(
            jobs_table.update()
            .where(jobs_table.c.run_id == RUN_ID)
            .values(parsing_mode="exploratory")
        )
    return ClipScribeWriterDB(engine=engine)


def test_transcript_sheet_shows_full_speech_even_without_a_digest(client, tmp_path):
    import io

    from openpyxl import load_workbook

    writer = _make_exploratory(tmp_path)
    writer.save_exploration_result(
        RUN_ID, 1, _explore_report(transcript=None), "gpt-test"
    )
    wb = load_workbook(
        io.BytesIO(client.get(f"/runs/{RUN_ID}/exploration/export").content)
    )
    rows = list(wb["Transcript"].iter_rows(min_row=2, values_only=True))
    assert rows == [("Full transcript", None, "hi")]


def test_job_exploration_export(client, tmp_path):
    import io

    from openpyxl import load_workbook

    # ABCD job: not available.
    assert client.get("/jobs/j1/exploration/export").status_code == 409

    writer = _make_exploratory(tmp_path)
    # No report yet: nothing to export.
    assert client.get("/jobs/j1/exploration/export").status_code == 409

    writer.save_exploration_result(RUN_ID, 1, _explore_report(), "gpt-test")

    resp = client.get("/jobs/j1/exploration/export?format=xlsx")
    assert resp.status_code == 200
    assert "_exploration.xlsx" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    # A cross-video summary, then one sheet per video (named after it).
    assert wb.sheetnames == ["Summary", "1. r1"]
    summary = list(wb["Summary"].iter_rows(min_row=2, values_only=True))
    assert summary[0][0] == "r1"
    assert summary[0][1:4] == (3.0, 2, 20.0)
    assert summary[0][-1] == "an ad"

    # The video sheet stacks the run's sections, transcript included.
    sheet = [row for row in wb["1. r1"].iter_rows(values_only=True)]
    titles = [
        row[0]
        for row in sheet
        if row[0] in {"Overview", "Chapters", "Transcript", "On-screen text", "Visuals"}
    ]
    assert titles == ["Overview", "Chapters", "Transcript", "On-screen text", "Visuals"]
    assert ("Full transcript", None, "hi") in [row[:3] for row in sheet]

    body = client.get("/jobs/j1/exploration/export?format=json").json()
    assert body["runs"][0]["video_name"] == "r1"
    assert body["runs"][0]["report"]["summary"] == "an ad"
    assert client.get("/jobs/j1/exploration/export?format=pdf").status_code == 400
